from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.opc.constants import RELATIONSHIP_TYPE as RT
from pathlib import Path
D=Document(); s=D.sections[0]
s.page_width=Inches(8.27); s.page_height=Inches(11.69)
s.top_margin=Inches(.55); s.bottom_margin=Inches(.55); s.left_margin=Inches(.65); s.right_margin=Inches(.65)
for name in ['Normal','Title','Subtitle','Heading 1','Heading 2','List Bullet']:
 st=D.styles[name]; st.font.name='Calibri'; st.font.color.rgb=RGBColor(0,0,0)
 st.paragraph_format.space_after=Pt(3)
 st.paragraph_format.line_spacing=1.03
D.styles['Normal'].font.size=Pt(10)
D.styles['Title'].font.size=Pt(24); D.styles['Title'].font.bold=True
D.styles['Heading 1'].font.size=Pt(11); D.styles['Heading 1'].font.bold=True
D.styles['Heading 1'].paragraph_format.space_before=Pt(9)
D.styles['Heading 1'].paragraph_format.space_after=Pt(4)
D.styles['List Bullet'].font.size=Pt(10)
D.styles['List Bullet'].paragraph_format.left_indent=Inches(.13)
D.styles['List Bullet'].paragraph_format.first_line_indent=Inches(-.13)
def p(t='',bold=False):
 x=D.add_paragraph(); x.add_run(t).bold=bold; return x
def head(t): D.add_paragraph(t,'Heading 1')
def bullet(t): D.add_paragraph(t,'List Bullet')
def link(p,label,url):
 h=OxmlElement('w:hyperlink'); h.set(qn('r:id'),p.part.relate_to(url,RT.HYPERLINK,is_external=True))
 r=OxmlElement('w:r'); prop=OxmlElement('w:rPr'); c=OxmlElement('w:color'); c.set(qn('w:val'),'222222'); prop.append(c); r.append(prop)
 t=OxmlElement('w:t'); t.text=label; r.append(t); h.append(r); p._p.append(h)
D.add_paragraph('NourEddin Abu Salhieh','Title')
p('Data Science and Artificial Intelligence',True)
x=p('Amman, Jordan  |  +962 79 621 7414  |  '); link(x,'noureddinabusalhieh@gmail.com','mailto:noureddinabusalhieh@gmail.com')
x=p(); link(x,'github.com/NourEddin004','https://github.com/NourEddin004'); x.add_run('  |  '); link(x,'LinkedIn','https://www.linkedin.com/in/noureddin-abu-salhieh-1135342aa/'); x.add_run('  |  '); link(x,'Portfolio','https://noureddin004.github.io/Protfolio/')
head('PROFILE')
p('Data Science and Artificial Intelligence degree candidate with hands-on experience in AI application development, web development and UI/UX design. Built Besign, an award-winning custom tailoring platform with on-device body measurement and AI garment design. Seeking an entry-level AI or data science role to apply Python, machine learning and problem-solving skills.')
head('EDUCATION')
p('Bachelor’s Degree in Data Science and Artificial Intelligence',True)
p('Applied Sciences Private University  |  Oct 2022 - Expected Sep 2026  |  GPA: 3.28')
head('TECHNICAL SKILLS')
for a,b in [('Programming','Python, Java, C++, JavaScript, PHP, HTML5, CSS3'),('Data and AI','pandas, NumPy, scikit-learn, PyTorch, TensorFlow, OpenCV, data analysis'),('Development and design','MySQL, Git, Flutter, Supabase, Figma, wireframing, UI/UX design')]:
 x=p(); x.add_run(a+': ').bold=True; x.add_run(b)
head('EXPERIENCE')
p('Full-Stack Web Development Intern | Orange Digital Center',True)
p('Jul 2026 - Present')
bullet('Built responsive web pages using HTML, CSS, Bootstrap and JavaScript, including form handling, local storage and Fetch API integration with Firebase Firestore.')
bullet('Developed PHP forms and a MySQL-backed charity platform with user accounts, sessions and an approval workflow; published training work in GitHub repositories.')
head('SELECTED PROJECTS')
p('Besign | AI-Powered Custom Tailoring Platform | Mar - Jun 2026',True)
bullet('Built a Flutter marketplace with on-device body measurement and AI garment design, using Dart, Riverpod, Supabase and TensorFlow Lite.')
bullet('Won 2nd place in Business & Social Technologies at the 13th National Technology Parade (Jul 2026), and Best Entrepreneurial Project at Applied Sciences Private University.')
p('HopeBridge | Charity Platform | Orange Digital Center',True)
bullet('Took the platform from user research, journeys and Figma wireframes to 28 HTML/CSS pages for donors, families and administrators; designed the database ER diagram.')
p('Invite Sites | Bilingual Invitation Websites',True)
bullet('Built personalised invitation sites with RSVP, maps and photos using HTML, CSS and JavaScript, supporting Arabic and English layouts.')
head('SELECTED CERTIFICATIONS AND TRAINING')
p('Microsoft Azure AI Fundamentals (AI-900) | Microsoft | May 2026')
p('Intermediate Python | DataCamp | Feb 2026')
p('HCIA-AI V4.0 Course | Huawei ICT Academy | Oct 2025')
head('LEADERSHIP AND LANGUAGES')
p('English Language Olympics | Volunteer and Organizer | 2023 - 2025',True)
p('Managed examination and judging halls, coordinated team presentations and oversaw contestant answer uploads.')
p('Arabic: Native  |  English: Excellent')
for el in list(D.styles.element.iter(qn('w:pBdr'))): el.getparent().remove(el)
for el in list(D.element.iter(qn('w:pBdr'))): el.getparent().remove(el)
D.core_properties.title='NourEddin Abu Salhieh CV'; D.core_properties.author='NourEddin Abu Salhieh'
D.save('output/cv/NourEddin-Abu-Salhieh-CV.docx')
