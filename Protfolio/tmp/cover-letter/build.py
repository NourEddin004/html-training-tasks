from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
from docx.opc.constants import RELATIONSHIP_TYPE as RT
D=Document(); s=D.sections[0]
s.page_width=Inches(8.27); s.page_height=Inches(11.69)
s.top_margin=Inches(.8); s.bottom_margin=Inches(.8); s.left_margin=Inches(.85); s.right_margin=Inches(.85)
for name in ['Normal','Title','Subtitle']:
 st=D.styles[name]; st.font.name='Calibri'; st.font.color.rgb=RGBColor(0,0,0)
 st.paragraph_format.line_spacing=1.12; st.paragraph_format.space_after=Pt(12)
D.styles['Normal'].font.size=Pt(11)
D.styles['Title'].font.size=Pt(23); D.styles['Title'].font.bold=True
D.styles['Title'].paragraph_format.space_after=Pt(7)
for e in list(D.styles.element.iter(qn('w:pBdr'))): e.getparent().remove(e)
def para(t,after=12):
 p=D.add_paragraph(t); p.paragraph_format.space_after=Pt(after); return p
def link(p,label,url):
 h=OxmlElement('w:hyperlink'); h.set(qn('r:id'),p.part.relate_to(url,RT.HYPERLINK,is_external=True))
 r=OxmlElement('w:r'); t=OxmlElement('w:t'); t.text=label; r.append(t); h.append(r); p._p.append(h)
D.add_paragraph('NourEddin Abu Salhieh','Title')
p=para('Amman, Jordan  |  +962 79 621 7414',3)
p=para('',3); link(p,'noureddinabusalhieh@gmail.com','mailto:noureddinabusalhieh@gmail.com')
p=para('',22); link(p,'Portfolio','https://noureddin004.github.io/Protfolio/'); p.add_run('  |  '); link(p,'LinkedIn','https://www.linkedin.com/in/noureddin-abu-salhieh-1135342aa/')
para('Dear Hiring Manager,',16)
para('I am seeking an entry-level opportunity in artificial intelligence or data science where I can contribute my programming, application development and UI/UX design skills. I am completing my bachelor’s degree in Data Science and Artificial Intelligence at Applied Sciences Private University, with graduation expected in September 2026. I am particularly interested in work that turns AI capabilities into useful applications.')
para('For my graduation project, Besign, I developed an AI-powered custom tailoring platform with on-device body measurement and AI garment design. The project brought together Flutter, Supabase and TensorFlow Lite, giving me practical experience connecting AI features with an application that people can use. Besign received second place in the Business & Social Technologies category at the 13th National Technology Parade and the university’s Best Entrepreneurial Project award.')
para('As a Full Stack Web Developer Intern at Orange Coding Academy, I have built responsive pages with HTML, CSS, Bootstrap and JavaScript, including form handling and Firebase Firestore integration. I also developed PHP forms and a MySQL-backed charity platform with user accounts, sessions and an approval workflow. This experience has strengthened my ability to work across the interface and server side of an application and maintain my work in GitHub repositories.')
para('Alongside my technical work, I have developed skills in Figma, user research and prototyping. My volunteer and organizer roles at the English Language Olympics involved managing examination and judging halls, coordinating presentations and overseeing answer uploads. These responsibilities gave me experience communicating clearly and coordinating tasks with others.')
para('I would welcome the opportunity to discuss how my project experience, technical foundation and interest in AI could support your team. Thank you for considering my application. My portfolio provides further examples of my work.',18)
para('Kind regards,',4)
p=para('NourEddin Abu Salhieh'); p.runs[0].bold=True
D.core_properties.title='NourEddin Abu Salhieh Cover Letter'; D.core_properties.author='NourEddin Abu Salhieh'
D.save('output/cover-letter/NourEddin-Abu-Salhieh-Cover-Letter.docx')
